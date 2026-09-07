"""
Извлечение текста из вложений задания (PDF / PPTX / DOCX / TXT).

Раньше такие файлы просто скачивались (services/google_service.
download_material_files) и на этом всё — их содержимое никак не попадало
в промпт Gemini. Из-за этого задание с прикреплённой презентацией или
документом (а не картинкой) уходило в Gemini практически пустым: только
заголовок и (если было) короткое описание — бот честно отвечал "не
прикрепили текст или изображение", хотя вложение по факту было.

Эти функции достают текст из файла, чтобы приложить его к промпту так
же, как контекст из учебников (см. gemini_service.find_relevant_book_context).
"""

from __future__ import annotations

import logging
from pathlib import Path
from uuid import uuid4

logger = logging.getLogger(__name__)

# Ограничение на длину извлечённого текста с одного файла — чтобы не
# раздувать промпт до неприличных размеров, если попадётся файл на
# сотню страниц.
_MAX_CHARS_PER_FILE = 12000

DOCUMENT_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt"}


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        slide_lines.append(row_text)
        if slide_lines:
            parts.append(f"[Слайд {i}]\n" + "\n".join(slide_lines))
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts)


def _extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".pptx": _extract_pptx,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
}


def extract_text(path: Path) -> str:
    """Возвращает извлечённый текст файла (обрезанный) либо '' при ошибке
    или неподдерживаемом формате."""
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return ""

    try:
        text = extractor(path).strip()
    except Exception:  # noqa: BLE001
        logger.exception("Не удалось извлечь текст из вложения %s", path)
        return ""

    if len(text) > _MAX_CHARS_PER_FILE:
        text = text[:_MAX_CHARS_PER_FILE] + "\n…(обрезано, файл длиннее)"
    return text


def extract_all(paths: list[Path]) -> str:
    """Извлекает текст из всех подходящих файлов и склеивает с заголовками
    по имени файла. Файлы, из которых ничего не извлеклось, пропускаются."""
    blocks = []
    for path in paths:
        text = extract_text(path)
        if text:
            blocks.append(f"--- Файл: {path.name} ---\n{text}")
    return "\n\n".join(blocks)


# Ограничение на число страниц PDF, которые рендерим в картинки —
# защита от случая, если кто-то прикрепит учебник на 300 страниц вместо
# одного задания.
_MAX_PDF_PAGES_AS_IMAGES = 6


def pdf_pages_as_images(path: Path, dest_dir: Path | None = None) -> list[Path]:
    """
    Рендерит страницы PDF в PNG-картинки.

    Раньше "невидимый" PDF (например, отсканированное фото задания без
    текстового слоя — так часто присылают домашку в школе) молча давал
    extract_text() == "" — и задание уходило в Gemini вообще без условия,
    хотя файл был. Теперь для таких сканов вместо пустого текста
    подставляются картинки страниц — Gemini прекрасно читает текст
    задания с изображения, так же как с обычного фото.

    Возвращает [] при любой ошибке (например, PyMuPDF не установлен) —
    вызывающий код должен быть готов к пустому списку и в таком случае
    сообщить пользователю, что файл прочитать не удалось.
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.warning(
            "PyMuPDF (пакет 'pymupdf') не установлен — не могу отрендерить "
            "страницы PDF %s как картинки для сканов без текстового слоя.",
            path,
        )
        return []

    out_dir = dest_dir or path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    images: list[Path] = []
    try:
        doc = fitz.open(str(path))
        try:
            for i, page in enumerate(doc):
                if i >= _MAX_PDF_PAGES_AS_IMAGES:
                    break
                pix = page.get_pixmap(dpi=150)
                img_path = out_dir / f"{path.stem}_p{i + 1}_{uuid4().hex[:6]}.png"
                pix.save(str(img_path))
                images.append(img_path)
        finally:
            doc.close()
    except Exception:  # noqa: BLE001
        logger.exception("Не удалось отрендерить страницы PDF %s как картинки", path)
        return []

    return images
